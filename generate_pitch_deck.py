#!/usr/bin/env python3
"""
TT Clinics - Bedrijfsplan PPTX Generator (Wit Thema)
Genereert een professionele 16:9 presentatie in het lichte/witte corporate thema van TT Clinics,
volledig uitgerust met officiële huisstijl-iconen, logo's en typografie.
"""

import sys
sys.path.insert(0, '/Users/matthias/Library/Python/3.9/lib/python/site-packages')
import os
import shutil
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# Initialiseer Presentatie (16:9 Widescreen)
prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# Kleurenpalet TT Clinics - WIT THEMA (Crisp Modern White & TT Orange)
C_BG = RGBColor(255, 255, 255)                # Puur wit
C_CANVAS = RGBColor(248, 250, 252)            # Zeer subtiel licht grijs/blauw voor zachte contrasten
C_CARD = RGBColor(255, 255, 255)              # Witte card
C_CARD_SOFT = RGBColor(248, 249, 251)         # Zachte card achtergrond
C_CARD_BORDER = RGBColor(226, 232, 240)       # Subtiele moderne rand (Slate 200)
C_CARD_BORDER_STRONG = RGBColor(203, 213, 225)
C_CARD_HIGHLIGHT_BG = RGBColor(255, 248, 242) # Warme oranje-witte gloed voor Company Battle (Flagship)
C_CARD_HIGHLIGHT_BORDER = RGBColor(255, 107, 0)

C_ORANGE = RGBColor(255, 107, 0)              # TT Oranje
C_ORANGE_DARK = RGBColor(217, 85, 0)          # Dieper oranje voor tekst op wit (hoog contrast)
C_ORANGE_LIGHT = RGBColor(255, 138, 51)
C_ORANGE_BADGE_BG = RGBColor(255, 243, 235)

C_TEXT_TITLE = RGBColor(17, 24, 39)           # Donker grafiet/zwart voor titels (Slate 900)
C_TEXT_HEAD = RGBColor(30, 41, 59)            # Slate 800 voor subtitels en koppen
C_TEXT_BODY = RGBColor(71, 85, 105)           # Slate 600 voor leestekst (contrastrijk & leesbaar)
C_TEXT_MUTED = RGBColor(148, 163, 184)        # Slate 400 voor voetnoten

FONT_HEAD = "Arial Black"
FONT_BODY = "Arial"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
IMG_LOGO = os.path.join(BASE_DIR, "images/logo-transparent.png")
IMG_PHOTO = os.path.join(BASE_DIR, "images/hero-photo.jpg")
IMG_QR = os.path.join(BASE_DIR, "images/qr-ttclinics.png")
IMG_STICKER = os.path.join(BASE_DIR, "images/logo-sticker.png")

# Officiële Huisstijl Iconen
ICON_TROPHY = os.path.join(BASE_DIR, "images/icon-trophy.png")
ICON_HEARTBEAT = os.path.join(BASE_DIR, "images/icon-heartbeat.png")
ICON_LOCATION = os.path.join(BASE_DIR, "images/icon-location.png")
ICON_SHIELD = os.path.join(BASE_DIR, "images/icon-shield.png")
ICON_CHART = os.path.join(BASE_DIR, "images/icon-chart.png")

def set_slide_background(slide, use_canvas=False):
    """Zet een strakke witte achtergrond met de kenmerkende TT Oranje accentbalk onderaan."""
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5)
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = C_CANVAS if use_canvas else C_BG
    bg_shape.line.fill.background()
    
    # Oranje accentlijn helemaal onderaan (TT Clinics handtekening)
    accent_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(7.42), Inches(13.333), Inches(0.08)
    )
    accent_bar.fill.solid()
    accent_bar.fill.fore_color.rgb = C_ORANGE
    accent_bar.line.fill.background()

def add_header(slide, category_text, title_text, slide_num=None):
    """Voegt een consistente professionele header toe aan de witte slide."""
    # Logo rechtsboven
    if os.path.exists(IMG_LOGO):
        slide.shapes.add_picture(IMG_LOGO, Inches(11.2), Inches(0.4), height=Inches(0.65))
    
    # Category tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(8), Inches(0.35))
    tf = tag_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = category_text.upper()
    p.font.name = FONT_HEAD
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = C_ORANGE_DARK
    
    # Slide Titel
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.65), Inches(10), Inches(0.8))
    tf2 = title_box.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.name = FONT_HEAD
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = C_TEXT_TITLE
    
    # Oranje scheidingslijn onder titel
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.48), Inches(4.5), Inches(0.04)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = C_ORANGE
    line.line.fill.background()
    
    # Footer tekst
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
    tf_f = footer_box.text_frame
    tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0
    pf = tf_f.paragraphs[0]
    num_str = f"  |  Slide {slide_num}" if slide_num else ""
    pf.text = f"TT Clinics  ·  Bedrijfsplan 2026-2028  ·  Edward Westhoff{num_str}"
    pf.font.name = FONT_BODY
    pf.font.size = Pt(8.5)
    pf.font.color.rgb = C_TEXT_MUTED

def add_card(slide, left, top, width, height, bg_color=C_CARD, border_color=C_CARD_BORDER, top_accent=None):
    """Maakt een strakke rounded card box in het witte thema."""
    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
        
    if top_accent:
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, left, top, width, Inches(0.08)
        )
        bar.fill.solid()
        bar.fill.fore_color.rgb = top_accent
        bar.line.fill.background()
        
    return card


# ==========================================
# SLIDE 1: COVER / TITELSLIDE (WIT THEMA)
# ==========================================
slide1 = prs.slides.add_slide(blank_layout)
set_slide_background(slide1)

# Subtiele warme oranje gloed rechtsboven
glow = slide1.shapes.add_shape(
    MSO_SHAPE.OVAL, Inches(8.5), Inches(-1.5), Inches(6), Inches(6)
)
glow.fill.solid()
glow.fill.fore_color.rgb = RGBColor(255, 245, 238)
glow.line.fill.background()

# Logo groot linksboven
if os.path.exists(IMG_LOGO):
    slide1.shapes.add_picture(IMG_LOGO, Inches(1.2), Inches(1.2), height=Inches(1.3))

# Badge
badge = slide1.shapes.add_shape(
    MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(2.9), Inches(3.6), Inches(0.4)
)
badge.fill.solid()
badge.fill.fore_color.rgb = C_ORANGE_BADGE_BG
badge.line.color.rgb = C_ORANGE
badge.line.width = Pt(1)
tf_b = badge.text_frame
tf_b.margin_top = Inches(0.06)
p = tf_b.paragraphs[0]
p.text = "★ BEDRIJFSPLAN 2026 - 2028"
p.font.name = FONT_HEAD
p.font.size = Pt(9.5)
p.font.bold = True
p.font.color.rgb = C_ORANGE_DARK
p.alignment = PP_ALIGN.CENTER

# Hoofdtitel
title_box = slide1.shapes.add_textbox(Inches(1.2), Inches(3.45), Inches(10), Inches(1.8))
tf = title_box.text_frame
tf.word_wrap = True
tf.margin_left = tf.margin_top = 0
p = tf.paragraphs[0]
p.text = "DYNAMIEK & SNELHEID"
p.font.name = FONT_HEAD
p.font.size = Pt(40)
p.font.bold = True
p.font.color.rgb = C_TEXT_TITLE

p2 = tf.add_paragraph()
p2.text = "Het sportieve bedrijfsuitje & vitaliteit op de werkvloer"
p2.font.name = FONT_BODY
p2.font.size = Pt(20)
p2.font.bold = True
p2.font.color.rgb = C_ORANGE
p2.space_before = Pt(8)

# Details blok
det_box = slide1.shapes.add_textbox(Inches(1.2), Inches(5.4), Inches(8), Inches(1.4))
tf_d = det_box.text_frame
tf_d.margin_left = 0
p = tf_d.paragraphs[0]
p.text = "Oprichter & Hoofdtrainer: Edward Westhoff"
p.font.name = FONT_HEAD
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = C_TEXT_HEAD

p2 = tf_d.add_paragraph()
p2.text = "TT Clinics  |  info@ttclinics.nl  |  www.ttclinics.nl  |  Heel Nederland"
p2.font.name = FONT_BODY
p2.font.size = Pt(11)
p2.font.color.rgb = C_TEXT_MUTED
p2.space_before = Pt(4)

# Huisstijl Trofee Icoon rechtsonder decoratief
if os.path.exists(ICON_TROPHY):
    slide1.shapes.add_picture(ICON_TROPHY, Inches(10.2), Inches(3.8), width=Inches(2.0))


# ==========================================
# SLIDE 2: EXECUTIVE SUMMARY
# ==========================================
slide2 = prs.slides.add_slide(blank_layout)
set_slide_background(slide2, use_canvas=True)
add_header(slide2, "Samenvatting", "Executive Summary: De Kracht van TT Clinics", 2)

cards_data = [
    ("HET CONCEPT", "Laagdrempelige, energieke tafeltennisclinics en bedrijfsuitjes voor MKB en corporate teams. Binnen 5 minuten ontstaat er plezier, gezonde rivaliteit en maximale teamdynamiek.", ICON_TROPHY),
    ("DE MISSIE", "Teams verbinden, beeldschermmoeheid doorbreken en vitaliteit bevorderen met 'schaken op topsnelheid'. Tafeltennis is fysiek én mentaal de ultieme breinfitness.", ICON_HEARTBEAT),
    ("DE OORSPRONG", "Wat begon als een bescheiden, gezellig tafeltennisclubje op de woensdagmorgen, is uitgegroeid tot een professionele partner voor bedrijfsvitaliteit en teambuilding.", ICON_LOCATION),
    ("DE AMBITIE", "Binnen 3 jaar uitgroeien tot de toonaangevende vitaliteitspartner in Nederland voor actieve bedrijfsevenementen met 100+ clinics per jaar en vaste partnerships.", ICON_CHART),
]

for i, (head, text, icon_path) in enumerate(cards_data):
    x = Inches(0.8 + (i % 2) * 5.95)
    y = Inches(1.8 + (i // 2) * 2.5)
    add_card(slide2, x, y, Inches(5.75), Inches(2.25), bg_color=C_CARD, border_color=C_CARD_BORDER, top_accent=C_ORANGE)
    
    # Brand icon in de kaart
    if os.path.exists(icon_path):
        slide2.shapes.add_picture(icon_path, x + Inches(4.85), y + Inches(0.25), width=Inches(0.65))
        
    tb = slide2.shapes.add_textbox(x + Inches(0.3), y + Inches(0.25), Inches(4.45), Inches(1.75))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = head
    p.font.name = FONT_HEAD
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = C_ORANGE_DARK
    
    p2 = tf.add_paragraph()
    p2.text = text
    p2.font.name = FONT_BODY
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = C_TEXT_BODY
    p2.space_before = Pt(8)


# ==========================================
# SLIDE 3: PROBLEEM & MARKTKANS
# ==========================================
slide3 = prs.slides.add_slide(blank_layout)
set_slide_background(slide3, use_canvas=True)
add_header(slide3, "Marktanalyse", "Het Probleem in Bedrijfsleven & Onze Marktkans", 3)

# Linker kolom: Het Probleem
add_card(slide3, Inches(0.8), Inches(1.8), Inches(5.75), Inches(5.0), bg_color=RGBColor(255, 250, 250), border_color=RGBColor(254, 202, 202), top_accent=RGBColor(220, 38, 38))

if os.path.exists(ICON_SHIELD):
    slide3.shapes.add_picture(ICON_SHIELD, Inches(5.8), Inches(1.95), width=Inches(0.6))

tb_p = slide3.shapes.add_textbox(Inches(1.1), Inches(2.05), Inches(5.15), Inches(4.5))
tf_p = tb_p.text_frame
tf_p.word_wrap = True
tf_p.margin_left = tf_p.margin_top = 0
p = tf_p.paragraphs[0]
p.text = "HET PROBLEEM OP KANTOOR"
p.font.name = FONT_HEAD
p.font.size = Pt(13)
p.font.color.rgb = RGBColor(185, 28, 28)

probs = [
    ("Schermmoeheid & Zitgedrag", "Medewerkers zitten gemiddeld 8-10 uur per dag achter een beeldscherm. Energielevels dalen en verzuim stijgt."),
    ("Vervreemding door Hybride Werken", "Collega's zien elkaar vaker via Zoom/Teams dan in het echt. Echte spontane verbinding en collegialiteit verdwijnt."),
    ("Saaie & Passieve Bedrijfsuitjes", "Standaard borrels of ongemakkelijke workshops zorgen zelden voor blijvende energie of échte interactie."),
    ("Hoge Drempels bij Sportuitjes", "Karten, padel of zeskamp sluit vaak collega's uit door conditieniveau, leeftijd of blessuregevoeligheid."),
]
for title, desc in probs:
    p_t = tf_p.add_paragraph()
    p_t.text = f"✕  {title}"
    p_t.font.name = FONT_HEAD
    p_t.font.size = Pt(10.5)
    p_t.font.color.rgb = C_TEXT_HEAD
    p_t.space_before = Pt(8)
    
    p_d = tf_p.add_paragraph()
    p_d.text = desc
    p_d.font.name = FONT_BODY
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = C_TEXT_BODY
    p_d.space_before = Pt(2)

# Rechter kolom: De Kans
add_card(slide3, Inches(6.78), Inches(1.8), Inches(5.75), Inches(5.0), bg_color=C_CARD_HIGHLIGHT_BG, border_color=C_CARD_HIGHLIGHT_BORDER, top_accent=C_ORANGE)

if os.path.exists(ICON_TROPHY):
    slide3.shapes.add_picture(ICON_TROPHY, Inches(11.75), Inches(1.95), width=Inches(0.6))

tb_k = slide3.shapes.add_textbox(Inches(7.08), Inches(2.05), Inches(5.15), Inches(4.5))
tf_k = tb_k.text_frame
tf_k.word_wrap = True
tf_k.margin_left = tf_k.margin_top = 0
p = tf_k.paragraphs[0]
p.text = "DE KANS VOOR TT CLINICS"
p.font.name = FONT_HEAD
p.font.size = Pt(13)
p.font.color.rgb = C_ORANGE_DARK

chances = [
    ("Groeiend Vitaliteitsbudget (WKR)", "Werkgevers investeren fors in preventieve vitaliteit, mentale fitheid en teambuilding met belastingvoordelen."),
    ("De Tafeltennis Herwaardering", "Elk modern kantoor heeft of wil een tafeltennistafel, maar benut deze zelden optimaal voor gestructureerde team clinics."),
    ("Schaken op Topsnelheid", "Tafeltennis triggert bewezen neuroplasticiteit: snelle reflexen, focus en mentale ontlading binnen 5 minuten."),
    ("Flexibel & Schaalbaar Model", "Wij komen met mobiele wedstrijduitrusting naar de bedrijfskantine óf organiseren het in een sfeervolle zaal incl. borrel."),
]
for title, desc in chances:
    p_t = tf_k.add_paragraph()
    p_t.text = f"✓  {title}"
    p_t.font.name = FONT_HEAD
    p_t.font.size = Pt(10.5)
    p_t.font.color.rgb = C_TEXT_HEAD
    p_t.space_before = Pt(8)
    
    p_d = tf_k.add_paragraph()
    p_d.text = desc
    p_d.font.name = FONT_BODY
    p_d.font.size = Pt(9.5)
    p_d.font.color.rgb = C_TEXT_BODY
    p_d.space_before = Pt(2)


# ==========================================
# SLIDE 4: DIENSTENAANBOD & PAKKETTEN
# ==========================================
slide4 = prs.slides.add_slide(blank_layout)
set_slide_background(slide4, use_canvas=True)
add_header(slide4, "Product & Propositie", "Het Dienstenaanbod: Drie Krachtige Formules", 4)

packages = [
    ("Team Kick-off", "€395", "1,5 uur · Tot 12 personen", ICON_LOCATION, [
        "Masterclass basistechniek & effectspin",
        "King of the Court toernooitje",
        "Professionele batjes & ballen inbegrepen",
        "Op eigen kantoor of sportlocatie",
        "Ideaal voor afdelingen & kleine teams"
    ], C_CARD, C_CARD_BORDER, None),
    
    ("Company Battle (Flagship)", "€745", "2,5 uur · Tot 12 personen", ICON_TROPHY, [
        "Complete clinic + trickshots van Edward",
        "Volledig verzorgd bedrijfscompetitie toernooi",
        "Officiële TT Clinics wisseltrofee & medailles",
        "Meest gekozen bedrijfsuitje & kick-off",
        "Optioneel: borrel & bittergarnituur arrangement"
    ], C_CARD_HIGHLIGHT_BG, C_CARD_HIGHLIGHT_BORDER, C_ORANGE),
    
    ("Vitaliteit & Maatwerk", "Op Maat", "Vanaf 12+ pers. / Meerdaags", ICON_HEARTBEAT, [
        "Gezondheids- en vitaliteitsweken op kantoor",
        "Meerdere trainers & mobiele wedstrijdtafels",
        "Bedrijfscompetitie opzet voor het hele jaar",
        "Volledig afgestemd op corporate programma",
        "Grote organisaties & evenementen"
    ], C_CARD, C_CARD_BORDER, None),
]

for i, (name, price, sub, icon_file, items, bg, border, accent) in enumerate(packages):
    x = Inches(0.8 + i * 3.98)
    y = Inches(1.8)
    add_card(slide4, x, y, Inches(3.78), Inches(5.0), bg_color=bg, border_color=border, top_accent=accent)
    
    # Huisstijl Icoon Badge bovenaan
    if os.path.exists(icon_file):
        slide4.shapes.add_picture(icon_file, x + Inches(3.78 - 0.95), y + Inches(0.2), width=Inches(0.75))
        
    tb = slide4.shapes.add_textbox(x + Inches(0.25), y + Inches(0.25), Inches(3.28), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    
    p = tf.paragraphs[0]
    p.text = name.upper()
    p.font.name = FONT_HEAD
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = C_ORANGE_DARK if accent else C_TEXT_HEAD
    
    p_pr = tf.add_paragraph()
    p_pr.text = price
    p_pr.font.name = FONT_HEAD
    p_pr.font.size = Pt(26)
    p_pr.font.bold = True
    p_pr.font.color.rgb = C_ORANGE if accent else C_TEXT_TITLE
    p_pr.space_before = Pt(4)
    
    p_s = tf.add_paragraph()
    p_s.text = sub
    p_s.font.name = FONT_BODY
    p_s.font.size = Pt(9.5)
    p_s.font.color.rgb = C_TEXT_MUTED
    p_s.space_before = Pt(2)
    
    # Scheidslijn
    p_div = tf.add_paragraph()
    p_div.text = "—" * 22
    p_div.font.size = Pt(8)
    p_div.font.color.rgb = C_CARD_BORDER_STRONG
    p_div.space_before = Pt(6)
    
    for it in items:
        p_i = tf.add_paragraph()
        p_i.text = f"✓ {it}"
        p_i.font.name = FONT_BODY
        p_i.font.size = Pt(9.5)
        p_i.font.color.rgb = C_TEXT_BODY
        p_i.space_before = Pt(6)


# ==========================================
# SLIDE 5: WAAROM TAFELTENNIS? (USP's)
# ==========================================
slide5 = prs.slides.add_slide(blank_layout)
set_slide_background(slide5, use_canvas=True)
add_header(slide5, "De Psychologie & Sport", "Waarom Tafeltennis? De Ultieme Teamformule", 5)

pillars = [
    ("100% Inclusief & Gelijkwaardig", "Geen fysiek overwicht of conditievereiste: stagiair en CEO staan direct op gelijk niveau tegenover elkaar aan tafel.", ICON_LOCATION),
    ("Schaken op Topsnelheid", "Activeert reflexen, hand-oogcoördinatie en breinplasticiteit. De ideale doorbreking van langdurige kantoor- en schermfocus.", ICON_HEARTBEAT),
    ("Direct Lachen & Rivaliteit", "Binnen 2 minuten ontstaan er spannende rally's, spectaculaire punten en hilarische missers. Gegarandeerd positieve teamenergie.", ICON_TROPHY),
    ("Volledig Zonder Blessures", "Veilige binnensport zonder fysiek contact. Iedereen kan zorgeloos meedoen in normale vrijetijdskleding of sportoutfit.", ICON_SHIELD),
]

for i, (title, text, icon_f) in enumerate(pillars):
    x = Inches(0.8 + (i % 2) * 5.95)
    y = Inches(1.8 + (i // 2) * 2.5)
    add_card(slide5, x, y, Inches(5.75), Inches(2.25), bg_color=C_CARD, border_color=C_CARD_BORDER, top_accent=C_ORANGE)
    
    if os.path.exists(icon_f):
        slide5.shapes.add_picture(icon_f, x + Inches(0.3), y + Inches(0.3), width=Inches(0.75))
        
    tb = slide5.shapes.add_textbox(x + Inches(1.2), y + Inches(0.25), Inches(4.3), Inches(1.75))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = title.upper()
    p.font.name = FONT_HEAD
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_HEAD
    
    p2 = tf.add_paragraph()
    p2.text = text
    p2.font.name = FONT_BODY
    p2.font.size = Pt(10.5)
    p2.font.color.rgb = C_TEXT_BODY
    p2.space_before = Pt(6)


# ==========================================
# SLIDE 6: DOELGROEPEN & KLANTPROFIEL
# ==========================================
slide6 = prs.slides.add_slide(blank_layout)
set_slide_background(slide6, use_canvas=True)
add_header(slide6, "Doelgroepsegmentatie", "Voor Wie? Vier Primaire Klantsegmenten", 6)

segments = [
    ("MKB & IT/Tech Bedrijven", "10 - 75 medewerkers", [
        "Hoge schermtijd, langdurig zittend werk",
        "Jong & dynamisch team dat competitie waardeert",
        "Zoeken informele, actieve vrijdagmiddag of kwartaalafsluiting",
        "Vaak al een tafeltennistafel aanwezig op kantoor"
    ]),
    ("Corporates & Zorg/Financieel", "Afdelingen van 12 - 50 pers.", [
        "Focus op vitaliteitsweken en duurzame inzetbaarheid",
        "Teambuilding na reorganisaties of fusies",
        "WKR-budget beschikbaar voor teamontwikkeling",
        "Boeken vaak 'Company Battle' of maatwerk"
    ]),
    ("Evenementenbureaus & Locaties", "B2B Partners", [
        "Zoeken een originele, interactieve break voor heidagen",
        "TT Clinics als vaste sportieve partner/module",
        "White-label of co-branded evenementen",
        "Terugkerende boekingen via intermediairs"
    ]),
    ("Ondernemersclubs & Netwerken", "20 - 60 ondernemers", [
        "Informele netwerkavond met toernooielement",
        "Lachen en netwerken in één dynamisch format",
        "Ideale sponsormogelijkheden rondom de trofee",
        "Directe leadgeneratie voor nieuwe bedrijfsuitjes"
    ]),
]

for i, (title, sub, bullet_list) in enumerate(segments):
    x = Inches(0.8 + (i % 2) * 5.95)
    y = Inches(1.8 + (i // 2) * 2.5)
    add_card(slide6, x, y, Inches(5.75), Inches(2.25), bg_color=C_CARD, border_color=C_CARD_BORDER)
    
    tb = slide6.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), Inches(5.15), Inches(1.85))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = title.upper()
    p.font.name = FONT_HEAD
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_ORANGE_DARK
    
    p_sub = tf.add_paragraph()
    p_sub.text = sub
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(9.5)
    p_sub.font.color.rgb = C_TEXT_MUTED
    p_sub.space_before = Pt(1)
    
    for b in bullet_list:
        p_b = tf.add_paragraph()
        p_b.text = f"• {b}"
        p_b.font.name = FONT_BODY
        p_b.font.size = Pt(9)
        p_b.font.color.rgb = C_TEXT_BODY
        p_b.space_before = Pt(3)


# ==========================================
# SLIDE 7: VERDIENMODEL & INKOMSTENSTROMEN
# ==========================================
slide7 = prs.slides.add_slide(blank_layout)
set_slide_background(slide7, use_canvas=True)
add_header(slide7, "Financieel Model", "Verdienmodel: Drie Inkomstenstromen", 7)

revenue_streams = [
    ("1. DIRECTE CLINIC OMZET", "De Kernmotor", ICON_TROPHY, [
        "Vaste pakketprijzen (€395 voor Kick-off, €745 voor Company Battle).",
        "Hoge brutomarge (80-85%): materiaal is reeds aangeschaft, hoofdkost is tijd & reiskosten.",
        "Opschaling: 2e trainer toevoegen bij groepen > 12 personen (+€250-€350 opslag)."
    ]),
    ("2. UPSELLS & ARRANGEMENTEN", "Hogere Orderwaarde", ICON_LOCATION, [
        "Borrel- & bittergarnituur arrangement (+€25 tot €45 p.p.).",
        "Gepersonaliseerde TT Clinics batjes met bedrijfslogo als blijvend aandenken.",
        "Professionele aftermovie / fotoreportage van het toernooi voor interne communicatie."
    ]),
    ("3. RECURRING COMPETITIES", "Terugkerende Omzet", ICON_CHART, [
        "Jaarlijkse interne kantoorcompetitie: 4 kwartaalrondes met finale dag.",
        "Maandelijkse vitaliteitsclinic als vast onderdeel van het bedrijfsprogramma.",
        "Voorspelbare cashflow en langdurige klantrelaties met vaste bedrijven."
    ]),
]

for i, (title, sub, icon_f, pts) in enumerate(revenue_streams):
    x = Inches(0.8 + i * 3.98)
    y = Inches(1.8)
    add_card(slide7, x, y, Inches(3.78), Inches(5.0), bg_color=C_CARD, border_color=C_CARD_BORDER, top_accent=C_ORANGE)
    
    if os.path.exists(icon_f):
        slide7.shapes.add_picture(icon_f, x + Inches(3.78 - 0.9), y + Inches(0.2), width=Inches(0.7))
        
    tb = slide7.shapes.add_textbox(x + Inches(0.25), y + Inches(0.25), Inches(3.28), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_HEAD
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = C_ORANGE_DARK
    
    p_s = tf.add_paragraph()
    p_s.text = sub
    p_s.font.name = FONT_BODY
    p_s.font.size = Pt(9.5)
    p_s.font.color.rgb = C_TEXT_MUTED
    p_s.space_before = Pt(2)
    
    p_d = tf.add_paragraph()
    p_d.text = "—" * 22
    p_d.font.size = Pt(8)
    p_d.font.color.rgb = C_CARD_BORDER_STRONG
    p_d.space_before = Pt(6)
    
    for pt in pts:
        p_pt = tf.add_paragraph()
        p_pt.text = f"• {pt}"
        p_pt.font.name = FONT_BODY
        p_pt.font.size = Pt(9.5)
        p_pt.font.color.rgb = C_TEXT_BODY
        p_pt.space_before = Pt(8)


# ==========================================
# SLIDE 8: MARKETING & GO-TO-MARKET STRATEGIE
# ==========================================
slide8 = prs.slides.add_slide(blank_layout)
set_slide_background(slide8, use_canvas=True)
add_header(slide8, "Commerciële Strategie", "Marketing & Go-To-Market: Hoe Werven We Klanten?", 8)

channels = [
    ("Website & Lokale B2B SEO", "ttclinics.nl", [
        "Hoog scoren op zoektermen als 'bedrijfsuitje tafeltennis', 'sportieve teamclinic', 'vitaliteitsdag kantoor'.",
        "Heldere pakketprijzen en directe online offerte-aanvraag flow.",
        "Social proof: reviews, foto's en video-impressies van gespeelde toernooien."
    ]),
    ("LinkedIn Content Marketing", "B2B Autoriteit", [
        "Korte dynamische video-reels van clinics: sfeer, tricks en juichende collega's.",
        "Artikelen over schermmoeheid, micro-pauzes en de voordelen van tafeltennis voor de hersenen.",
        "Direct benaderen van HR-managers, Office Managers en Event Coördinatoren."
    ]),
    ("Vliegwiel / Mond-tot-Mond", "Elke Clinic is een Pitch", [
        "Deelnemers aan een clinic zijn managers of medewerkers bij andere projecten/bedrijven.",
        "Iedere clinic levert direct nieuwe ambassadeurs en doorverwijzingen op.",
        "Visitekaartjes met directe QR-code uitdelen na afloop van de clinic."
    ]),
    ("Partnerships & Lokale Clubs", "Netwerk", [
        "Samenwerking met tafeltennisverenigingen voor zaalverhuur en bardiensten.",
        "Koppeling met cateringbedrijven en evenementenlocaties voor complete dagarrangementen.",
        "Actieve aanwezigheid bij lokale business clubs en netwerklunches."
    ]),
]

for i, (title, sub, items) in enumerate(channels):
    x = Inches(0.8 + (i % 2) * 5.95)
    y = Inches(1.8 + (i // 2) * 2.5)
    add_card(slide8, x, y, Inches(5.75), Inches(2.25), bg_color=C_CARD, border_color=C_CARD_BORDER)
    
    tb = slide8.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), Inches(5.15), Inches(1.85))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = title.upper()
    p.font.name = FONT_HEAD
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_TEXT_HEAD
    
    p_s = tf.add_paragraph()
    p_s.text = sub
    p_s.font.name = FONT_BODY
    p_s.font.size = Pt(9.5)
    p_s.font.color.rgb = C_ORANGE_DARK
    p_s.space_before = Pt(1)
    
    for it in items:
        p_i = tf.add_paragraph()
        p_i.text = f"• {it}"
        p_i.font.name = FONT_BODY
        p_i.font.size = Pt(9)
        p_i.font.color.rgb = C_TEXT_BODY
        p_i.space_before = Pt(3)


# ==========================================
# SLIDE 9: OPERATIE, UITRUSTING & LOGISTIEK
# ==========================================
slide9 = prs.slides.add_slide(blank_layout)
set_slide_background(slide9, use_canvas=True)
add_header(slide9, "Operatie & Uitvoering", "Operationele Uitrusting & Schaalbare Logistiek", 9)

ops = [
    ("Professionele Uitrusting", "Wedstrijdklasse Materiaal", [
        "Mobiele ITTF-goedgekeurde wedstrijdtafels en mobiele netsets.",
        "Hoogwaardige trainingsbatjes met professionele rubbers (maximale controle en spin).",
        "Officiële 3-sterren wedstrijdballen met TT Clinics logobedrukking.",
        "Gepersonaliseerde wisseltrofeeën en toernooimedailles."
    ]),
    ("Locatie Flexibiliteit", "Overal Speelklaar", [
        "Kantoor/Bedrijfskantine: Transformatie in 20 minuten tot toernooiarena.",
        "Sportaccommodaties: Vaste afspraken met zalen inclusief kleedkamers en afsluitende borrel.",
        "Outdoor opties: Bij mooi weer op buitenpleinen of dakterrassen."
    ]),
    ("Schaalbaar Trainersmodel", "Capaciteit Uitbreiden", [
        "Edward Westhoff als hoofdtrainer en hét vertrouwde gezicht van TT Clinics.",
        "Pool van ervaren competitiespelers en gediplomeerde trainers voor piekmomenten.",
        "Standaard clinic-draaiboek garandeert consistente kwaliteit en enthousiasme."
    ]),
    ("Klantreis & Ontzorging", "Van Aanvraag tot Trofee", [
        "Binnen 24 uur persoonlijk contact en offerte op maat.",
        "Voorbereiding: Inventarisatie van ruimte en specifieke wensen.",
        "Uitvoering: Complete toernooileiding, muziek, coaching en prijsuitreiking.",
        "Nazorg: Foto's toesturen en evaluatiegesprek voor vervolgclinic."
    ]),
]

for i, (title, sub, items) in enumerate(ops):
    x = Inches(0.8 + (i % 2) * 5.95)
    y = Inches(1.8 + (i // 2) * 2.5)
    add_card(slide9, x, y, Inches(5.75), Inches(2.25), bg_color=C_CARD, border_color=C_CARD_BORDER)
    
    tb = slide9.shapes.add_textbox(x + Inches(0.3), y + Inches(0.2), Inches(5.15), Inches(1.85))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    p = tf.paragraphs[0]
    p.text = title.upper()
    p.font.name = FONT_HEAD
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = C_ORANGE_DARK
    
    p_s = tf.add_paragraph()
    p_s.text = sub
    p_s.font.name = FONT_BODY
    p_s.font.size = Pt(9.5)
    p_s.font.color.rgb = C_TEXT_MUTED
    p_s.space_before = Pt(1)
    
    for it in items:
        p_i = tf.add_paragraph()
        p_i.text = f"• {it}"
        p_i.font.name = FONT_BODY
        p_i.font.size = Pt(9)
        p_i.font.color.rgb = C_TEXT_BODY
        p_i.space_before = Pt(3)


# ==========================================
# SLIDE 10: FINANCIËLE ROADMAP 2026 - 2028
# ==========================================
slide10 = prs.slides.add_slide(blank_layout)
set_slide_background(slide10, use_canvas=True)
add_header(slide10, "Groeipad & Doelen", "Financiële Roadmap & Mijlpalen 2026 - 2028", 10)

milestones = [
    ("FASE 1: 2026", "Fundament & Tractie", "€25.000 - €35.000", ICON_LOCATION, [
        "40 - 50 Clinics verzorgen.",
        "Conversiegerichte website & leadmachine live.",
        "Eerste 5 vaste corporate accounts binnenhalen.",
        "Focus op regio Randstad / Midden-Nederland.",
        "Klanttevredenheidsscore > 9.0."
    ], C_CARD, C_CARD_BORDER, None),
    ("FASE 2: 2027", "Opschaling & Abonnementen", "€60.000 - €75.000", ICON_CHART, [
        "80 - 100 Clinics per jaar.",
        "Uitrol terugkerende bedrijfscompetities.",
        "Aanstellen van 2e trainer voor piekperiodes.",
        "Partnerships met 3 landelijke evenementenbureaus.",
        "Introductie custom batjes met bedrijfslogo."
    ], C_CARD_HIGHLIGHT_BG, C_CARD_HIGHLIGHT_BORDER, C_ORANGE),
    ("FASE 3: 2028", "Landelijke Autoriteit", "€100.000+", ICON_TROPHY, [
        "150+ Clinics & toernooien per jaar.",
        "Vaste vitaliteitspartner voor 20+ corporate klanten.",
        "Landelijke dekking met gecertificeerde trainers.",
        "Eigen TT Clinics Bedrijfscompetitie Trophy Finale.",
        "Gevestigde A-merkpositie in bedrijfsvitaliteit."
    ], C_CARD, C_CARD_BORDER, None),
]

for i, (phase, title, rev, icon_f, pts, bg, border, accent) in enumerate(milestones):
    x = Inches(0.8 + i * 3.98)
    y = Inches(1.8)
    add_card(slide10, x, y, Inches(3.78), Inches(5.0), bg_color=bg, border_color=border, top_accent=accent)
    
    if os.path.exists(icon_f):
        slide10.shapes.add_picture(icon_f, x + Inches(3.78 - 0.9), y + Inches(0.2), width=Inches(0.7))
        
    tb = slide10.shapes.add_textbox(x + Inches(0.25), y + Inches(0.25), Inches(3.28), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    
    p = tf.paragraphs[0]
    p.text = phase
    p.font.name = FONT_HEAD
    p.font.size = Pt(11.5)
    p.font.bold = True
    p.font.color.rgb = C_ORANGE_DARK if accent else C_TEXT_HEAD
    
    p_t = tf.add_paragraph()
    p_t.text = title
    p_t.font.name = FONT_HEAD
    p_t.font.size = Pt(13)
    p_t.font.bold = True
    p_t.font.color.rgb = C_TEXT_TITLE
    p_t.space_before = Pt(2)
    
    p_rev = tf.add_paragraph()
    p_rev.text = f"Omzetdoel: {rev}"
    p_rev.font.name = FONT_HEAD
    p_rev.font.size = Pt(12)
    p_rev.font.bold = True
    p_rev.font.color.rgb = C_ORANGE
    p_rev.space_before = Pt(4)
    
    p_d = tf.add_paragraph()
    p_d.text = "—" * 22
    p_d.font.size = Pt(8)
    p_d.font.color.rgb = C_CARD_BORDER_STRONG
    p_d.space_before = Pt(6)
    
    for pt in pts:
        p_pt = tf.add_paragraph()
        p_pt.text = f"✓ {pt}"
        p_pt.font.name = FONT_BODY
        p_pt.font.size = Pt(9.5)
        p_pt.font.color.rgb = C_TEXT_BODY
        p_pt.space_before = Pt(6)


# ==========================================
# SLIDE 11: CONCLUSIE & CONTACT
# ==========================================
slide11 = prs.slides.add_slide(blank_layout)
set_slide_background(slide11, use_canvas=True)
add_header(slide11, "Conclusie & Samenwerking", "Klaar Voor Meer Energie Op De Werkvloer?", 11)

# Linker kaart: Trainer & Verhaal
add_card(slide11, Inches(0.8), Inches(1.8), Inches(7.5), Inches(5.0), bg_color=C_CARD, border_color=C_CARD_BORDER, top_accent=C_ORANGE)

# Foto Edward
if os.path.exists(IMG_PHOTO):
    slide11.shapes.add_picture(IMG_PHOTO, Inches(1.1), Inches(2.2), width=Inches(2.2))

tb_c = slide11.shapes.add_textbox(Inches(3.6), Inches(2.1), Inches(4.4), Inches(4.3))
tf_c = tb_c.text_frame
tf_c.word_wrap = True
tf_c.margin_left = tf_c.margin_top = 0

p = tf_c.paragraphs[0]
p.text = "EDWARD WESTHOFF"
p.font.name = FONT_HEAD
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = C_TEXT_TITLE

p_sub = tf_c.add_paragraph()
p_sub.text = "Oprichter & Hoofdtrainer · TT Clinics"
p_sub.font.name = FONT_HEAD
p_sub.font.size = Pt(10)
p_sub.font.bold = True
p_sub.font.color.rgb = C_ORANGE_DARK
p_sub.space_before = Pt(2)

p_quote = tf_c.add_paragraph()
p_quote.text = (
    "\"TT Clinics is het bewijs dat de mooiste dingen soms heel klein beginnen. "
    "Wat startte als een bescheiden clubje op de woensdagmorgen, is inmiddels uitgegroeid "
    "tot een professionele partner voor bedrijfsvitaliteit en teambuilding. "
    "Tafeltennis is dé manier om laagdrempelig in beweging te komen en als team onvergetelijk plezier te maken.\""
)
p_quote.font.name = FONT_BODY
p_quote.font.size = Pt(10)
p_quote.font.italic = True
p_quote.font.color.rgb = C_TEXT_BODY
p_quote.space_before = Pt(10)

p_exp = tf_c.add_paragraph()
p_exp.text = "★ 15+ Jaren Ervaring    ★ 500+ Spelers Gecoacht    ★ Heel Nederland"
p_exp.font.name = FONT_HEAD
p_exp.font.size = Pt(9.5)
p_exp.font.bold = True
p_exp.font.color.rgb = C_ORANGE_DARK
p_exp.space_before = Pt(12)


# Rechter kaart: Direct Contact & QR
add_card(slide11, Inches(8.55), Inches(1.8), Inches(3.98), Inches(5.0), bg_color=C_CARD, border_color=C_CARD_BORDER, top_accent=C_ORANGE)
tb_info = slide11.shapes.add_textbox(Inches(8.85), Inches(2.1), Inches(3.38), Inches(4.3))
tf_info = tb_info.text_frame
tf_info.word_wrap = True
tf_info.margin_left = tf_info.margin_top = 0

p = tf_info.paragraphs[0]
p.text = "DIRECT CONTACT"
p.font.name = FONT_HEAD
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = C_TEXT_TITLE

info_items = [
    ("E-mail:", "info@ttclinics.nl"),
    ("Telefoon:", "06 - 12345678"),
    ("Website:", "www.ttclinics.nl"),
    ("Locaties:", "Op kantoor & sportzalen"),
]

for label, val in info_items:
    p_i = tf_info.add_paragraph()
    p_i.text = f"{label} {val}"
    p_i.font.name = FONT_BODY
    p_i.font.size = Pt(10)
    p_i.font.color.rgb = C_TEXT_BODY
    p_i.space_before = Pt(6)

# QR code afbeelding
if os.path.exists(IMG_QR):
    slide11.shapes.add_picture(IMG_QR, Inches(9.8), Inches(4.9), width=Inches(1.5))
    qr_caption = slide11.shapes.add_textbox(Inches(8.85), Inches(6.5), Inches(3.38), Inches(0.4))
    tf_q = qr_caption.text_frame
    p_q = tf_q.paragraphs[0]
    p_q.text = "Scan voor website & team clinics"
    p_q.font.name = FONT_BODY
    p_q.font.size = Pt(8.5)
    p_q.font.color.rgb = C_TEXT_MUTED
    p_q.alignment = PP_ALIGN.CENTER


# ==========================================
# OPSLAAN VAN HET BESTAND
# ==========================================
out_dir = BASE_DIR
output_file = os.path.join(out_dir, "TT_Clinics_Bedrijfsplan.pptx")
prs.save(output_file)
print(f"Presentatie succesvol gegenereerd: {output_file}")

# Kopieer ook direct naar ~/Downloads/TT_Clinics_Bedrijfsplan.pptx voor snel openen door de gebruiker
downloads_dir = "/Users/matthias/Downloads"
if os.path.exists(downloads_dir):
    user_copy = os.path.join(downloads_dir, "TT_Clinics_Bedrijfsplan.pptx")
    shutil.copy2(output_file, user_copy)
    print(f"Kopie geplaatst in Downloads: {user_copy}")
